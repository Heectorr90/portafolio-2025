from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("artifacts/Hector_Ramirez_CV_ATS.docx")


PROFILE = [
    "Desarrollador Full Stack orientado a la creación de aplicaciones web, APIs REST y automatización de procesos. Experiencia con Vue.js, Angular, Laravel, PHP, MySQL, JavaScript y herramientas modernas de despliegue. Enfocado en construir soluciones mantenibles, responsivas y útiles para optimizar procesos de negocio.",
]

TECHNICAL_SKILLS = {
    "Frontend": "Vue.js, Angular, React, JavaScript, TypeScript, HTML5, CSS3, Vuetify, Tailwind CSS, Responsive Design",
    "Backend": "PHP, Laravel, REST APIs, Node.js, Python, Django, Livewire, JWT, autenticación, validaciones",
    "Bases de datos": "MySQL, modelado relacional, consultas SQL, integridad de datos",
    "DevOps y herramientas": "Git, GitHub, Docker, Railway, Hostinger, Postman, VS Code, Codex, despliegue web",
    "Integraciones": "PayPal API, Mercado Pago API, OpenAI, Gemini, servicios SMTP/transaccionales",
}

EXPERIENCE = [
    {
        "title": "Desarrollo de Software",
        "company": 'Bit Electronics "Centro de servicio"',
        "period": "2022 - 2025",
        "bullets": [
            "Automatización, monitoreo y resolución de incidencias técnicas mediante scripts y herramientas backend.",
            "Desarrollo de soluciones internas para reducir tareas repetitivas y mejorar la productividad operativa.",
            "Soporte a procesos técnicos con enfoque en mantenimiento, mejora continua y estabilidad de sistemas.",
        ],
    },
    {
        "title": "Programador Front-end",
        "company": "CTED Corporativo Tecnológico y Educativo S.A. de C.V.",
        "period": "2021 - 2022",
        "bullets": [
            "Participación en desarrollo y mantenimiento de aplicaciones web con HTML, CSS, JavaScript y Angular.",
            "Construcción de interfaces responsivas y apoyo en la mejora de componentes visuales existentes.",
            "Colaboración en tareas de implementación frontend durante etapa de formación profesional.",
        ],
    },
]

PROJECTS = [
    {
        "name": "E-commerce Full Stack",
        "stack": "Laravel 10, Angular 17, MySQL, JWT, PayPal API, Mercado Pago API, TypeScript",
        "bullets": [
            "Desarrollo de tienda online con flujo de usuarios, productos, carrito, pagos en línea y panel administrativo.",
            "Implementación de API REST para operación de ecommerce y separación entre backend y SPA frontend.",
        ],
    },
    {
        "name": "Sistema de Inventario de Equipos",
        "stack": "Laravel, MySQL, Livewire, Tailwind CSS",
        "bullets": [
            "Sistema web con estructura jerárquica de categorías, CRUD, búsqueda, filtros y paginación.",
            "Generación automática de códigos únicos de inventario con validaciones y formularios dinámicos.",
        ],
    },
    {
        "name": "Blog Personal con Roles",
        "stack": "Vue.js, Vuetify, Laravel, MySQL, JWT",
        "bullets": [
            "Aplicación full stack con autenticación, búsqueda, categorías y panel administrativo con roles de usuario.",
            "Integración frontend-backend mediante API y manejo de sesiones con JWT.",
        ],
    },
    {
        "name": "DashForge",
        "stack": "Vue 3, Pinia, Laravel, MySQL, Tailwind CSS",
        "bullets": [
            "Aplicación para crear dashboards dinámicos con widgets personalizables y drag and drop.",
            "Soporte para modo local y persistencia de datos con usuarios autenticados.",
        ],
    },
    {
        "name": "Chatbot de IA",
        "stack": "React, OpenAI, Gemini",
        "bullets": [
            "Chatbot inteligente para responder preguntas y asistir usuarios en tiempo real mediante IA generativa.",
        ],
    },
]

EDUCATION = [
    {
        "degree": "Ingeniería en Sistemas Computacionales",
        "institution": "Tecnológico Universitario de Toluca",
        "period": "2023 - 2024",
    }
]


def set_cell_text(cell, text, bold=False):
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(10)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color in [
        ("Heading 1", 12, "1F4D78"),
        ("Heading 2", 10.5, "1F4D78"),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)


def add_heading(doc, text):
    paragraph = doc.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text.upper())
    run.bold = True
    add_bottom_border(paragraph)


def add_bottom_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D9E2F3")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.first_line_indent = Inches(-0.15)
    run = paragraph.add_run(text)
    run.font.size = Pt(10.2)
    return paragraph


def add_role(doc, title, company, period, bullets):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run(f"{title} | {company} | {period}")
    run.bold = True
    run.font.size = Pt(10.5)
    for bullet in bullets:
        add_bullet(doc, bullet)


def add_project(doc, name, stack, bullets):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run(name)
    run.bold = True
    run.font.size = Pt(10.5)
    stack_run = paragraph.add_run(f" | {stack}")
    stack_run.italic = True
    stack_run.font.size = Pt(10)
    for bullet in bullets:
        add_bullet(doc, bullet)


def build_doc():
    doc = Document()
    configure_document(doc)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(0)
    run = name.add_run("HECTOR RAMIREZ")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor.from_string("0B2545")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("Desarrollador Full Stack | Laravel | Vue.js | Angular | APIs REST")
    run.bold = True
    run.font.size = Pt(10.5)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(6)
    contact.add_run("Toluca, Estado de Mexico | +52 722 571 9891 | developerbit035@gmail.com | ")
    add_hyperlink(contact, "GitHub", "https://github.com/Heectorr90")
    contact.add_run(" | ")
    add_hyperlink(contact, "LinkedIn", "https://www.linkedin.com/in/hector-ramirez-909577215/")

    add_heading(doc, "Perfil profesional")
    for paragraph_text in PROFILE:
        p = doc.add_paragraph(paragraph_text)
        p.paragraph_format.space_after = Pt(5)

    add_heading(doc, "Competencias técnicas")
    for label, value in TECHNICAL_SKILLS.items():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        label_run = p.add_run(f"{label}: ")
        label_run.bold = True
        p.add_run(value)

    add_heading(doc, "Experiencia profesional")
    for role in EXPERIENCE:
        add_role(doc, role["title"], role["company"], role["period"], role["bullets"])

    add_heading(doc, "Proyectos relevantes")
    for project in PROJECTS:
        add_project(doc, project["name"], project["stack"], project["bullets"])

    add_heading(doc, "Educación")
    for item in EDUCATION:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(f"{item['degree']} | {item['institution']} | {item['period']}")
        run.bold = True

    add_heading(doc, "Palabras clave ATS")
    keywords = (
        "Desarrollador Full Stack, Desarrollador Web, Frontend Developer, Backend Developer, "
        "Laravel, PHP, Vue.js, Angular, React, JavaScript, TypeScript, MySQL, REST API, API REST, "
        "JWT, Git, GitHub, Docker, Postman, Railway, Hostinger, Tailwind CSS, Livewire, Node.js, Python"
    )
    doc.add_paragraph(keywords)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_doc()
